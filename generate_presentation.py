"""Generation de la presentation ScrumFirst (DataOps) via python-pptx."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# Palette
SAND = RGBColor(247, 246, 240)
NAVY = RGBColor(15, 23, 42)
NAVY_SOFT = RGBColor(30, 41, 59)
ACCENT = RGBColor(71, 85, 105)

TITLE_FONT = "Trebuchet MS"
BODY_FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]


def set_bg(slide, color=SAND):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color=NAVY):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, font=BODY_FONT,
             bold=False, color=NAVY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6, line=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color
    return tb


def add_rich_paragraph(tf, segments, size=18, align=PP_ALIGN.LEFT,
                       space_after=10, line=1.2, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for text, bold in segments:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = BODY_FONT
        r.font.color.rgb = NAVY if bold else NAVY_SOFT
    return p


def slide_header(slide, index, title, subtitle=None):
    """Barre d'en-tete minimaliste commune."""
    add_rect(slide, Inches(0), Inches(0), Inches(0.18), SLIDE_H, NAVY)
    add_text(slide, Inches(0.7), Inches(0.45), Inches(1.6), Inches(0.4),
             f"0{index}", size=13, font=TITLE_FONT, bold=True, color=ACCENT)
    add_text(slide, Inches(0.7), Inches(0.85), Inches(11.9), Inches(1.0),
             title, size=32, font=TITLE_FONT, bold=True, color=NAVY)
    y = Inches(1.85)
    if subtitle:
        add_text(slide, Inches(0.72), Inches(1.7), Inches(11.5), Inches(0.5),
                 subtitle, size=15, font=BODY_FONT, color=ACCENT)
        y = Inches(2.3)
    add_rect(slide, Inches(0.72), y - Inches(0.12), Inches(2.2), Emu(28000), NAVY)
    return y + Inches(0.35)


def add_footer(slide, text="ScrumFirst  |  Direction Data & Intelligence Artificielle"):
    add_rect(slide, Inches(0), Inches(7.1), SLIDE_W, Emu(19000), NAVY)
    add_text(slide, Inches(0.72), Inches(7.12), Inches(9), Inches(0.3),
             text, size=10, font=BODY_FONT, color=ACCENT)


def add_bullets(slide, y, items, width=11.9, x=0.72, size=17, gap=0.32):
    """items = liste de (label, description) affichees en puces courtes."""
    tb = slide.shapes.add_textbox(Inches(x), y, Inches(width), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (label, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(14)
        p.line_spacing = 1.25
        b = p.add_run()
        b.text = "\u25aa  "
        b.font.size = Pt(size)
        b.font.bold = True
        b.font.name = BODY_FONT
        b.font.color.rgb = NAVY
        r = p.add_run()
        r.text = label
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = BODY_FONT
        r.font.color.rgb = NAVY
        if desc:
            d = p.add_run()
            d.text = " : " + desc
            d.font.size = Pt(size)
            d.font.bold = False
            d.font.name = BODY_FONT
            d.font.color.rgb = NAVY_SOFT
    return tb


# ---------------------------------------------------------------- DIAPOSITIVE 1
s = prs.slides.add_slide(BLANK)
set_bg(s)
add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.22), NAVY)
add_rect(s, Inches(0), Inches(7.28), SLIDE_W, Inches(0.22), NAVY)
add_rect(s, Inches(1.1), Inches(2.25), Inches(1.7), Emu(38000), NAVY)

add_text(s, Inches(1.1), Inches(1.55), Inches(6), Inches(0.5),
         "PRESENTATION EXECUTIVE", size=15, font=TITLE_FONT, bold=True,
         color=ACCENT)

add_text(s, Inches(1.1), Inches(2.6), Inches(11.1), Inches(2.0),
         "ScrumFirst : L'Agilité à l'Échelle\npour les Projets Data & IA",
         size=46, font=TITLE_FONT, bold=True, color=NAVY, line=1.1)

add_text(s, Inches(1.1), Inches(4.75), Inches(10.5), Inches(1.0),
         "Piloter la valeur continue, automatiser la gouvernance\net structurer le DataOps.",
         size=20, font=BODY_FONT, color=NAVY_SOFT, line=1.3)

add_rect(s, Inches(1.1), Inches(6.35), Inches(3.4), Emu(19000), NAVY)
add_text(s, Inches(1.1), Inches(6.5), Inches(11), Inches(0.5),
         "Présentation Exécutive - Direction Data & Intelligence Artificielle.",
         size=14, font=BODY_FONT, bold=True, color=ACCENT)

# ---------------------------------------------------------------- DIAPOSITIVE 2
s = prs.slides.add_slide(BLANK)
set_bg(s)
y = slide_header(s, 2, "Pourquoi ScrumFirst ?",
                 "Répondre aux défis du DataOps")
add_bullets(s, y, [
    ("Aligner la R&D IA et le Métier",
     "Fin des projets boîtes noires grâce à un Product Backlog unique axé sur un score de valeur métier."),
    ("Maîtriser la Qualité des Pipelines",
     "Fin de la dette technique via l'automatisation stricte et obligatoire des critères de qualité."),
    ("Traçabilité (Data Lineage)",
     "Historisation complète des cycles de développement, des blocages et des décisions d'équipe."),
])
add_footer(s)

# ---------------------------------------------------------------- DIAPOSITIVE 3
s = prs.slides.add_slide(BLANK)
set_bg(s)
y = slide_header(s, 3, "Sécurisation des Livrables",
                 "La Révolution de la Qualité : la Definition of Done (DoD) — Gouvernance automatisée")
add_bullets(s, y, [
    ("La Grille DoD Globale",
     "Critères de qualité non négociables (tests unitaires passants, versioning de modèle validé, non-régression)."),
    ("Blocage strict à la source",
     "Impossible de passer un item à « DONE » ou de le promouvoir en Incrément si la DoD est < 100 %."),
    ("Droit à l'erreur intégré",
     "Rejet d'un item de « DONE » vers le Backlog avec purge automatique des critères en cas d'anomalie en production."),
])
add_footer(s)

# ---------------------------------------------------------------- DIAPOSITIVE 4
s = prs.slides.add_slide(BLANK)
set_bg(s)
y = slide_header(s, 4, "Une Architecture Conçue pour le Passage à l'Échelle",
                 "Flexibilité Organisationnelle : le Multi-Équipes Parallèle")
add_bullets(s, y, [
    ("Source Unique de Vérité",
     "Un seul Product Backlog global pour le produit Data, géré par le Product Owner."),
    ("Sprints Parallèles et Étanches",
     "Plusieurs équipes (Team Data Engineering, Team Data Science) travaillent sur le même backlog."),
    ("Verrous de Sécurité Contextuels",
     "Algorithme interdisant à une même équipe d'ouvrir deux sprints actifs, sans bloquer les autres équipes."),
])
add_footer(s)

# ---------------------------------------------------------------- DIAPOSITIVE 5
s = prs.slides.add_slide(BLANK)
set_bg(s)
y = slide_header(s, 5, "Traçabilité Absolue des Rituels Agiles",
                 "La « Time Machine » du Projet : Auditabilité & Rétrospectives")
add_bullets(s, y, [
    ("Compteur et Suivi des Daily",
     "Suivi dynamique des Daily Scrums quotidiens (ex : X / 20) avec clôture définitive par le Scrum Master."),
    ("Typage des Notes Quotidiennes",
     "Historique catégorisé par nature d'événement (Avancement, Obstacle/Blocage, Annonce)."),
    ("Interface Épurée par Filtres",
     "Sélection instantanée [Produit ➔ Sprint ➔ Timeline] évitant la surcharge cognitive pour l'audit."),
])
add_footer(s)

# ---------------------------------------------------------------- DIAPOSITIVE 6
s = prs.slides.add_slide(BLANK)
set_bg(s)
y = slide_header(s, 6, "Un Moteur Relationnel Robuste et Performant",
                 "Architecture Technique & Modèle de Données (Prisma / Postgres)")
add_bullets(s, y, [
    ("Modélisation SQL Propre",
     "Tables interconnectées [Product ➔ BacklogItem ➔ Sprint ➔ Increment ➔ DailyNote]."),
    ("Optimisation Algorithmique",
     "Indexations complexes sur les clés stratégiques (teamId, productId, status) pour des temps de réponse instantanés sous forte charge."),
    ("Sécurité des Rôles",
     "Matrice de droits d'accès étanche codée en dur (Scrum Team en écriture / Stakeholders en lecture seule absolue)."),
])
add_footer(s)

OUT = "ScrumFirst_Presentation.pptx"
prs.save(OUT)
print(f"OK -> {OUT} ({len(prs.slides.slides if hasattr(prs.slides,'slides') else prs.slides._sldIdLst)} diapos)")
