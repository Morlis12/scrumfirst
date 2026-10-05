import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Initialisation de la présentation au format moderne 16:9
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Charte graphique de l'application (Sand-300 / Navy-900)
SAND_300 = RGBColor(247, 246, 240)
NAVY_900 = RGBColor(15, 23, 42)

def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = SAND_300

def add_slide_title(slide, text):
    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.833), Inches(1.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = 'Trebuchet MS'
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = NAVY_900

slide_layout = prs.slide_layouts[6] # Configuration slide vierge

# -------------------------------------------------------------
# DIAPOSITIVE 1 : SLIDE D'OUVERTURE
# -------------------------------------------------------------
slide1 = prs.slides.add_slide(slide_layout)
apply_background(slide1)

title_box = slide1.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(11.833), Inches(2.5))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "ScrumFirst"
p.font.name = 'Trebuchet MS'
p.font.size = Pt(54)
p.font.bold = True
p.font.color.rgb = NAVY_900

p2 = tf.add_paragraph()
p2.text = "L'Agilité Multi-Tenant à l'Échelle pour les Projets Data & IA"
p2.font.name = 'Calibri'
p2.font.size = Pt(24)
p2.font.color.rgb = NAVY_900
p2.space_before = Pt(10)

subtitle_box = slide1.shapes.add_textbox(Inches(0.75), Inches(5.5), Inches(11.833), Inches(1.0))
tf_sub = subtitle_box.text_frame
p_sub = tf_sub.paragraphs[0]
p_sub.text = "Présentation Exécutive — Direction Data & Intelligence Artificielle"
p_sub.font.name = 'Calibri'
p_sub.font.size = Pt(14)
p_sub.font.color.rgb = NAVY_900

# -------------------------------------------------------------
# DIAPOSITIVE 2 : LE POSITIONNEMENT DATAOPS
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(slide_layout)
apply_background(slide2)
add_slide_title(slide2, "Pourquoi ScrumFirst ? Répondre aux défis du DataOps")

content_box = slide2.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.833), Inches(4.5))
tf = content_box.text_frame
tf.word_wrap = True

points = [
    ("Aligner la R&D IA et la Valeur Métier", "Fin de l'effet boîte noire des algorithmes grâce à un Product Backlog unique et indépendant par produit, priorisé selon un score algorithmique de rentabilité (Valeur Métier / Effort)."),
    ("Gouvernance et Qualité des Pipelines", "Élimination active de la dette technique via l'automatisation stricte et l'application obligatoire de la Definition of Done (DoD) avant intégration des modèles."),
    ("Traçabilité et Auditabilité (Data Lineage)", "Historisation complète et centralisée des décisions techniques, des rituels et des blocages d'équipes pour répondre aux exigences réglementaires.")
]

for title, desc in points:
    p = tf.add_paragraph()
    p.text = f"• {title} :"
    p.font.name = 'Calibri'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_900
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.name = 'Calibri'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = NAVY_900
    p_desc.space_after = Pt(14)
    p_desc.level = 1

# -------------------------------------------------------------
# DIAPOSITIVE 3 : LA RÉVOLUTION DE LA QUALITÉ (DoD)
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(slide_layout)
apply_background(slide3)
add_slide_title(slide3, "Sécurisation des Livrables via la Definition of Done")

content_box = slide3.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.833), Inches(4.5))
tf = content_box.text_frame
tf.word_wrap = True

points3 = [
    ("Grille de Qualité Globale et Inflexible", "Chaque produit définit ses critères de qualité non négociables (ex: taux de couverture des tests unitaires, validation du schéma de données, non-régression)."),
    ("Blocage à la Source et Micro-Incréments", "Impossible de valider un item et de le promouvoir en Incrément si la jauge DoD est inférieure à 100%. Validation au fil de l'eau (CI/CD agile)."),
    ("Droit à l'erreur et Purge Automatique", "Option de rejet 'Retour au backlog depuis DONE' en cas d'anomalie détectée en production : nettoie instantanément les tables de liaison et réinitialise l'item.")
]

for title, desc in points3:
    p = tf.add_paragraph()
    p.text = f"• {title} :"
    p.font.name = 'Calibri'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_900
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.name = 'Calibri'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = NAVY_900
    p_desc.space_after = Pt(14)
    p_desc.level = 1

# -------------------------------------------------------------
# DIAPOSITIVE 4 : L'ARCHITECTURE MULTI-TENANT COMMERCIALE
# -------------------------------------------------------------
slide4 = prs.slides.add_slide(slide_layout)
apply_background(slide4)
add_slide_title(slide4, "Modèle commercial Multi-Tenant & Multi-Équipes")

content_box = slide4.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.833), Inches(4.5))
tf = content_box.text_frame
tf.word_wrap = True

points4 = [
    ("Workspace multi-locataire (SaaS commercial)", "Introduction du modèle Workspace. Isolement absolu des requêtes par identifiant unique. L'application démarre 100% vierge pour chaque organisation."),
    ("Gouvernance par Produit (Framework LeSS)", "Centralisation d'un Product Owner et d'un Scrum Master uniques au niveau du produit. Toutes les équipes créées sous ce produit héritent du binôme stratégique."),
    ("Contrainte RH native du Scrum Guide", "Chaque équipe accepte maximum 8 Développeurs/Data Scientists (10 membres au total avec PO/SM) pour garantir l'efficacité stricte du Daily de 15 minutes.")
]

for title, desc in points4:
    p = tf.add_paragraph()
    p.text = f"• {title} :"
    p.font.name = 'Calibri'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_900
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.name = 'Calibri'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = NAVY_900
    p_desc.space_after = Pt(14)
    p_desc.level = 1

# -------------------------------------------------------------
# DIAPOSITIVE 5 : SUIVI INTERACTIF & REVIEWS RELOCALISÉES
# -------------------------------------------------------------
slide5 = prs.slides.add_slide(slide_layout)
apply_background(slide5)
add_slide_title(slide5, "La 'Time Machine' du Projet : Suivi et Audits de Crise")

content_box = slide5.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.833), Inches(4.5))
tf = content_box.text_frame
tf.word_wrap = True

points5 = [
    ("Compteur de Répétition des Daily", "Gestion de la récurrence des Daily (X / TOTAL) avec bouton de clôture définitive libérant dynamiquement l'étape suivante (Sprint Review)."),
    ("Fil d'Actualité Temporel Typé", "Historisation des commentaires catégorisés par nature d'événement (Avancement, Obstacle/Blocage, Annonce) pour isoler instantanément les crises."),
    ("Relocalisation et Droits de la Sprint Review", "Panneau de Review collaboratif en bas de Kanban : zone de feedback ouverte uniquement aux Stakeholders (Métiers), Scrum Team en lecture seule.")
]

for title, desc in points5:
    p = tf.add_paragraph()
    p.text = f"• {title} :"
    p.font.name = 'Calibri'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_900
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.name = 'Calibri'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = NAVY_900
    p_desc.space_after = Pt(14)
    p_desc.level = 1

# -------------------------------------------------------------
# DIAPOSITIVE 6 : INFRASTRUCTURE TECHNIQUE (PRISMA/POSTGRES)
# -------------------------------------------------------------
slide6 = prs.slides.add_slide(slide_layout)
apply_background(slide6)
add_slide_title(slide6, "Infrastructure Technique & Industrialisation")

content_box = slide6.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.833), Inches(4.5))
tf = content_box.text_frame
tf.word_wrap = True

points6 = [
    ("Moteur Relationnel Robuste", "Modélisation SQL propre sous PostgreSQL via l'ORM Prisma. Tables interconnectées : Workspace, User, Product, Team, Sprint, Increment, DailyNote, RetrospectiveIdea."),
    ("Optimisation pour la Forte Charge", "Mise en place d'indexations complexes en base de données sur les couples stratégiques (teamId, productId, status) pour garantir un temps de réponse instantané."),
    ("Sécurité et Intégration Continue Vercel", "Déploiement découplé et automatisé. Validation totale de la compilation TypeScript (tsc 0 erreur) assurant une intégration continue (CI/CD) sans faille.")
]

for title, desc in points6:
    p = tf.add_paragraph()
    p.text = f"• {title} :"
    p.font.name = 'Calibri'
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY_900
    
    p_desc = tf.add_paragraph()
    p_desc.text = desc
    p_desc.font.name = 'Calibri'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = NAVY_900
    p_desc.space_after = Pt(14)
    p_desc.level = 1

prs.save('ScrumFirst_Presentation_DataOps.pptx')
print("Fichier PPTX généré avec succès sous le nom 'ScrumFirst_Presentation_DataOps.pptx'.")
